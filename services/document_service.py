import os
import uuid
import pymupdf  # PyMuPDF
from docx import Document as DocxDocument
from werkzeug.utils import secure_filename
from models.database_models import Document
from extensions import db

ALLOWED_EXTENSIONS = {'pdf', 'docx'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def validate_file(file):
    if file.filename == '':
        return False, "No selected file"
    if not allowed_file(file.filename):
        return False, "Invalid file type. Only PDF and DOCX are allowed."
    return True, None

def save_uploaded_file(file, user_id, upload_folder):
    """Saves the file to disk and creates an initial database record."""
    document_id = str(uuid.uuid4())
    original_filename = file.filename
    safe_filename = secure_filename(original_filename)
    
    # Extract file extension
    file_type = safe_filename.rsplit('.', 1)[1].lower()
    
    # Target directory: uploads/{user_id}/{document_id}/
    target_dir = os.path.join(upload_folder, str(user_id), document_id)
    
    # Path traversal validation check
    abs_upload_folder = os.path.abspath(upload_folder)
    abs_target_dir = os.path.abspath(target_dir)
    if not abs_target_dir.startswith(abs_upload_folder):
        raise ValueError("Path traversal attempt detected.")
        
    os.makedirs(target_dir, exist_ok=True)
    
    # Save the physical file
    file_path = os.path.join(target_dir, safe_filename)
    file.save(file_path)
    
    # Get file size
    file_size = os.path.getsize(file_path)
    
    # Create DB record
    new_doc = Document(
        user_id=user_id,
        filename=safe_filename,
        original_filename=original_filename,
        file_path=file_path,
        file_type=file_type,
        file_size=file_size,
        status='processing'
    )
    db.session.add(new_doc)
    db.session.commit()
    
    return new_doc

def clean_text(text):
    """Cleans extracted text to remove excessive whitespace."""
    if not text:
        return ""
    # Normalize whitespace
    return " ".join(text.split())

def extract_pdf_text(filepath):
    """Extracts text from a PDF file using PyMuPDF."""
    text_content = ""
    page_count = 0
    try:
        doc = pymupdf.open(filepath)
        page_count = len(doc)
        for page_num in range(page_count):
            page = doc.load_page(page_num)
            page_text = page.get_text("text")
            if page_text:
                text_content += f"\n--- Page {page_num + 1} ---\n{page_text}"
        doc.close()
        return True, clean_text(text_content), page_count
    except Exception as e:
        print(f"Error extracting PDF text: {e}")
        return False, "", 0

def extract_docx_text(filepath):
    """Extracts text from a DOCX file using python-docx."""
    text_content = ""
    try:
        doc = DocxDocument(filepath)
        for para in doc.paragraphs:
            if para.text.strip():
                text_content += para.text + "\n"
        # Word doesn't have reliable page counts in docx structure without rendering. 
        # We'll use a rough estimate (approx 500 words per page)
        words = len(text_content.split())
        page_count = max(1, words // 500)
        return True, clean_text(text_content), page_count
    except Exception as e:
        print(f"Error extracting DOCX text: {e}")
        return False, "", 0

def analyze_document_statistics(text, page_count):
    """Calculates document statistics."""
    word_count = len(text.split())
    character_count = len(text)
    return {
        "page_count": page_count,
        "word_count": word_count,
        "character_count": character_count
    }

def process_document(document_id):
    """Main pipeline for text extraction."""
    doc = db.session.get(Document, document_id)
    if not doc:
        return False
        
    try:
        # Extract text based on file type
        if doc.file_type == 'pdf':
            success, text, page_count = extract_pdf_text(doc.file_path)
        elif doc.file_type == 'docx':
            success, text, page_count = extract_docx_text(doc.file_path)
        else:
            success = False
            
        if success:
            stats = analyze_document_statistics(text, page_count)
            doc.extracted_text = text
            doc.page_count = stats["page_count"]
            doc.word_count = stats["word_count"]
            doc.character_count = stats["character_count"]
            
            # Index document for RAG
            from services.embedding_service import index_document
            index_success = index_document(str(doc.id), text, doc.user_id, doc.original_filename)
            if index_success:
                doc.status = 'completed'
            else:
                doc.status = 'indexed_failed'
                
        else:
            doc.status = 'failed'
            doc.extracted_text = "Text extraction failed or format unsupported."
            
        db.session.commit()
        return success
    except Exception as e:
        print(f"Document processing failed for {document_id}: {e}")
        doc.status = 'failed'
        db.session.commit()
        return False
