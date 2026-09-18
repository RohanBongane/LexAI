import docx

def create_docx():
    doc = docx.Document()
    doc.add_heading('LexAI - Complete Technical Audit', 0)
    
    with open('audit_report.txt', 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
            
    doc.save('LexAI_Technical_Audit.docx')
    print("DOCX created successfully.")

if __name__ == "__main__":
    create_docx()
